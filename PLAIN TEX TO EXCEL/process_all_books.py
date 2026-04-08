import json
import re
import os
import pandas as pd
from pathlib import Path

# ==========================================
# CONFIGURATION
# ==========================================
BOOKS_DIR = "books"
HIERARCHY_DIR = "hierarchy"
OUTPUT_DIR = "output"
XLSX_DIR = "File Sheet"

# Ensure output directories exist
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(XLSX_DIR, exist_ok=True)

# ==========================================
# HELPER FUNCTIONS
# ==========================================

def detect_arabic(text):
    """Check if text contains Arabic characters."""
    return bool(re.search(r'[\u0600-\u06FF]', text))

def process_content_tags(content):
    """Isolate Arabic text and wrap appropriately."""
    arabic_pattern = r'([\u0600-\u06FF][\u0600-\u06FF\s\d،؛؟.,\-\(\)\[\]«»\'"{}]+[\u0600-\u06FF]|[\u0600-\u06FF]+)'

    def replacer(match):
        return f'</c>\n<a>{match.group(1)}</a>\n<c>'

    processed = re.sub(arabic_pattern, replacer, content)
    return f"<c>\n{processed}\n</c>"

def clean_empty_tags(text):
    """Remove empty <c></c> tags."""
    return re.sub(r'<c>\s*</c>', '', text, flags=re.DOTALL)

def parse_tagged_text(text):
    """Convert <c>...</c> and <a>...</a> tagged text to HTML."""
    import unicodedata

    def normalize_lines(raw):
        return [line.strip() for line in raw.split('\n') if line.strip()]

    def is_bengali_text(raw):
        return re.search(r'[\u0980-\u09FF]', raw) is not None

    def is_arabic_connector_fragment(raw):
        compact = ''.join(raw.split())
        if not compact or len(compact) > 12 or is_bengali_text(compact):
            return False
        honorific_chars = {'\ufdfa', '\u0610', '\u0611', '\u0612', '\u0613', '\u0614'}
        allowed_symbols = set(".:,;!?()[]{}\"'/-\\|<>`~@#$%^&*_+=\u00ab\u00bb\u061f\u060c\u061b")
        for ch in compact:
            category = unicodedata.category(ch)
            if category[0] in ('P', 'S'):
                continue
            if ch in honorific_chars or ch in allowed_symbols:
                continue
            return False
        return True

    def join_connector(prev_arabic, connector):
        connector = ' '.join(normalize_lines(connector))
        if not connector:
            return prev_arabic
        no_space_before = set(".,:;!?)]}\u00bb\u061f\u060c\u061b")
        if connector[0] in no_space_before:
            return prev_arabic.rstrip() + connector
        return prev_arabic.rstrip() + ' ' + connector

    raw_blocks = []
    parts = re.split(r'(<[ca]>.*?</[ca]>)', text, flags=re.DOTALL)
    for part in parts:
        part = part.strip()
        if not part:
            continue
        c_match = re.match(r'<c>(.*?)</c>', part, re.DOTALL)
        a_match = re.match(r'<a>(.*?)</a>', part, re.DOTALL)
        if c_match:
            raw_blocks.append(('c', c_match.group(1).strip()))
        elif a_match:
            raw_blocks.append(('a', a_match.group(1).strip()))

    merged_blocks = []
    i = 0
    while i < len(raw_blocks):
        tag, value = raw_blocks[i]
        if tag == 'c' and is_arabic_connector_fragment(value):
            prev_is_arabic = bool(merged_blocks and merged_blocks[-1][0] == 'a')
            next_is_arabic = bool(i + 1 < len(raw_blocks) and raw_blocks[i + 1][0] == 'a')
            if prev_is_arabic and next_is_arabic:
                left = join_connector(merged_blocks[-1][1], value)
                right = ' '.join(normalize_lines(raw_blocks[i + 1][1]))
                merged_blocks[-1] = ('a', f"{left} {right}".strip())
                i += 2
                continue
            if prev_is_arabic:
                merged_blocks[-1] = ('a', join_connector(merged_blocks[-1][1], value))
                i += 1
                continue
            if next_is_arabic:
                next_value = ' '.join(normalize_lines(raw_blocks[i + 1][1]))
                merged_blocks.append(('a', f"{value.strip()} {next_value}".strip()))
                i += 2
                continue
        if merged_blocks and merged_blocks[-1][0] == tag:
            merged_blocks[-1] = (tag, merged_blocks[-1][1] + '\n' + value)
        else:
            merged_blocks.append((tag, value))
        i += 1

    html_parts = []
    for tag, value in merged_blocks:
        paragraphs = normalize_lines(value)
        if not paragraphs:
            continue
        if tag == 'c':
            html_parts.append(f'<div class="bengali-text">{"<br>".join(paragraphs)}</div>')
        else:
            html_parts.append(f'<div class="arabic-text">{"<br>".join(paragraphs)}</div>')
    return '\n'.join(html_parts)

def build_content_html(content_list):
    """Build HTML from content list."""
    html = ""
    for item in content_list:
        page = item.get("page", "")
        if page:
            html += f'<div class="page-marker">{page}</div>\n'
        html += parse_tagged_text(item["text"]) + "\n"
    return html


def save_book_as_xlsx(book_data, book_name, output_path):
    """Save book data as an XLSX file with structured sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()

    # Remove default sheet
    wb.remove(wb.active)

    # Create Chapters sheet
    ws_chapters = wb.create_sheet("Chapters")
    ws_chapters.append(["chapter_id", "title"])
    for ch in book_data["chapters"]:
        ws_chapters.append([ch["id"], ch["title"]])

    # Create Sections sheet
    ws_sections = wb.create_sheet("Sections")
    ws_sections.append(["section_id", "chapter_id", "title"])
    for ch in book_data["chapters"]:
        for sec in ch.get("sections", []):
            ws_sections.append([sec["id"], ch["id"], sec["title"]])

    # Create Content sheet
    ws_content = wb.create_sheet("Content")
    ws_content.append(["chapter_id", "section_id", "page_number", "text"])

    for ch in book_data["chapters"]:
        # Intro content (no section)
        for item in ch.get("intro_content", []):
            ws_content.append([ch["id"], None, item.get("page", ""), item.get("text", "")])

        # Section content
        for sec in ch.get("sections", []):
            for item in sec.get("content", []):
                ws_content.append([ch["id"], sec["id"], item.get("page", ""), item.get("text", "")])

    # Style headers for all sheets
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4A6785", end_color="4A6785", fill_type="solid")

    for ws in [ws_chapters, ws_sections, ws_content]:
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, min(len(str(cell.value)), 50))
                except:
                    pass
            ws.column_dimensions[column_letter].width = max_length + 2

    wb.save(output_path)
    print(f"  Saved XLSX: {output_path}")

# ==========================================
# MAIN PROCESSING FUNCTION
# ==========================================

def process_book(book_name, txt_path, hierarchy_path):
    """Process a single book and return structured data."""
    print(f"  Processing: {book_name}")

    # Load files
    with open(txt_path, "r", encoding="utf-8") as f:
        text = f.read()

    with open(hierarchy_path, "r", encoding="utf-8") as f:
        hierarchy = json.load(f)

    # Build lookup lists
    chapters_list = [item["chapter"] for item in hierarchy]
    sections_list = [sec for item in hierarchy for sec in item["sections"]]

    def heading_replacer(match):
        raw_heading = match.group(1).strip()
        clean_heading = " ".join(raw_heading.split())
        if clean_heading in chapters_list:
            return f"<chapter>{clean_heading}</chapter>"
        else:
            return f"<section>{clean_heading}</section>"

    # Convert tags
    text = re.sub(r'<page_number>\s*(.*?)\s*</page_number>', r'<p>\1</p>', text, flags=re.DOTALL)
    text = re.sub(r'<heading>\s*(.*?)\s*</heading>', heading_replacer, text, flags=re.DOTALL)

    # Process content
    def process_content(match):
        content = match.group(1).strip()
        processed = process_content_tags(content)
        return processed

    text = re.sub(r'<text>(.*?)</text>', process_content, text, flags=re.DOTALL)
    text = clean_empty_tags(text)

    # Parse to structure
    from bs4 import BeautifulSoup
    xml_output = f"<book>\n{text}\n</book>"
    soup = BeautifulSoup(xml_output, "html.parser")

    # Build book data structure
    book_data = []
    current_chapter = None
    current_section = None
    current_page = None
    current_content_parts = []

    def flush_content():
        nonlocal current_content_parts
        if not current_content_parts and not current_section and not current_chapter:
            return

        merged = []
        for tag, txt in current_content_parts:
            if merged and merged[-1][0] == tag:
                merged[-1] = (tag, merged[-1][1] + '\n' + txt)
            else:
                merged.append((tag, txt))

        combined = '\n'.join(f'<{tag}>{txt}</{tag}>' for tag, txt in merged)

        if current_section:
            current_section["content"].append({
                "page": current_page or "",
                "text": combined
            })
        elif current_chapter:
            current_chapter["intro_content"].append({
                "page": current_page or "",
                "text": combined
            })
        current_content_parts = []

    chapter_id = 0
    section_id = 0

    for element in soup.book.children:
        if element.name is None:
            continue

        tag = element.name
        text_val = element.text.strip()

        if not text_val:
            continue

        if tag == 'p':
            current_page = text_val

        elif tag == 'chapter':
            flush_content()
            chapter_id += 1
            current_chapter = {
                "id": chapter_id,
                "title": text_val,
                "intro_content": [],
                "sections": []
            }
            current_section = None
            book_data.append(current_chapter)

        elif tag == 'section':
            flush_content()
            section_id += 1
            current_section = {
                "id": section_id,
                "title": text_val,
                "content": []
            }
            if current_chapter:
                current_chapter["sections"].append(current_section)

        elif tag in ['c', 'a']:
            current_content_parts.append((tag, text_val))

    flush_content()

    return book_data

def generate_multi_book_viewer(books_index):
    """Generate the multi-book HTML viewer that loads books dynamically."""

    # Build book selector options
    book_options = '\n'.join([
        f'<option value="{book["id"]}">{book["name"]}</option>'
        for book in books_index
    ])

    # Only store the index (small metadata) in the HTML
    index_json = json.dumps(books_index, ensure_ascii=False, indent=2)

    html = f'''<!DOCTYPE html>
<html lang="bn" dir="ltr">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Islamic Book Viewer</title>
<style>
  @import url('https://fonts.googleapis.com/css2?family=Noto+Sans+Bengali:wght@400;600;700&family=Noto+Naskh+Arabic:wght@400;600;700&display=swap');

  * {{ margin: 0; padding: 0; box-sizing: border-box; }}

  :root {{
    --page-bg: #faf6f0;
    --page-border: #d4c5a9;
    --text-dark: #2c1810;
    --text-muted: #6b5b4f;
    --accent: #8b4513;
    --accent-light: #d4a574;
    --arabic-bg: #f0ebe3;
    --arabic-border: #c9b896;
    --sidebar-bg: #3b2f24;
    --sidebar-text: #e8ddd0;
    --sidebar-hover: #d4a574;
    --shadow: 0 2px 20px rgba(0,0,0,0.1);
  }}

  body {{
    font-family: 'Noto Sans Bengali', 'Segoe UI', sans-serif;
    background: #e8e0d4;
    color: var(--text-dark);
    line-height: 1.9;
    display: flex;
    min-height: 100vh;
  }}

  /* ===== SIDEBAR / TOC ===== */
  .sidebar {{
    width: 320px;
    min-width: 320px;
    background: var(--sidebar-bg);
    color: var(--sidebar-text);
    padding: 20px;
    overflow: hidden;
    position: fixed;
    top: 0;
    left: 0;
    height: 100vh;
    z-index: 100;
    transition: transform 0.3s ease;
    display: flex;
    flex-direction: column;
  }}

  .sidebar-header {{
    text-align: center;
    padding-bottom: 15px;
    margin-bottom: 15px;
    border-bottom: 1px solid rgba(255,255,255,0.15);
    flex-shrink: 0;
  }}

  .sidebar-header h2 {{
    font-size: 1.2rem;
    font-weight: 700;
    color: var(--accent-light);
    margin-bottom: 8px;
  }}

  /* Book Selector */
  .book-selector {{
    width: 100%;
    padding: 10px 12px;
    font-size: 0.95rem;
    font-family: inherit;
    background: rgba(255,255,255,0.1);
    color: var(--sidebar-text);
    border: 1px solid rgba(255,255,255,0.2);
    border-radius: 8px;
    cursor: pointer;
    margin-bottom: 10px;
  }}

  .book-selector:focus {{
    outline: none;
    border-color: var(--accent-light);
  }}

  .book-selector option {{
    background: var(--sidebar-bg);
    color: var(--sidebar-text);
  }}

  .book-info {{
    font-size: 0.75rem;
    opacity: 0.6;
    margin-top: 5px;
  }}

  .toc {{
    list-style: none;
    margin-top: 15px;
    flex: 1;
    min-height: 0;
    overflow-y: auto;
    padding-right: 4px;
    padding-bottom: 20px;
  }}

  .toc-chapter {{
    margin-bottom: 8px;
  }}

  .toc-chapter > a {{
    display: block;
    padding: 8px 12px;
    color: var(--sidebar-text);
    text-decoration: none;
    font-weight: 600;
    font-size: 0.9rem;
    border-radius: 6px;
    transition: all 0.2s;
  }}

  .toc-chapter > a:hover,
  .toc-chapter > a.active {{
    background: rgba(255,255,255,0.1);
    color: var(--sidebar-hover);
  }}

  .toc-sections {{
    list-style: none;
    padding-left: 16px;
    max-height: 0;
    overflow: hidden;
    transition: max-height 0.3s ease;
  }}

  .toc-chapter.expanded .toc-sections {{
    max-height: 3000px;
  }}

  .toc-sections li a {{
    display: block;
    padding: 4px 12px;
    color: rgba(232, 221, 208, 0.7);
    text-decoration: none;
    font-size: 0.8rem;
    border-left: 2px solid transparent;
    transition: all 0.2s;
  }}

  .toc-sections li a:hover,
  .toc-sections li a.active {{
    color: var(--sidebar-hover);
    border-left-color: var(--sidebar-hover);
  }}

  .sidebar-footer {{
    flex-shrink: 0;
    padding-top: 12px;
    margin-top: 12px;
    border-top: 1px solid rgba(255,255,255,0.12);
  }}

  .sidebar-action {{
    display: flex;
    align-items: center;
    justify-content: center;
    width: 100%;
    padding: 10px 12px;
    border-radius: 10px;
    border: 1px solid rgba(255,255,255,0.2);
    background: rgba(255,255,255,0.08);
    color: var(--sidebar-text);
    font-family: inherit;
    font-size: 0.9rem;
    font-weight: 600;
    cursor: pointer;
    transition: all 0.2s ease;
  }}

  .sidebar-action:hover {{
    background: rgba(255,255,255,0.14);
    color: var(--sidebar-hover);
    border-color: rgba(212,165,116,0.7);
  }}

  .json-merger-modal {{
    display: none;
    position: fixed;
    inset: 0;
    z-index: 400;
    background: rgba(17, 12, 8, 0.6);
    align-items: center;
    justify-content: center;
    padding: 24px;
  }}

  .json-merger-modal.open {{
    display: flex;
  }}

  .json-merger-dialog {{
    width: min(1100px, 96vw);
    height: min(820px, 92vh);
    background: #fff;
    border-radius: 12px;
    overflow: hidden;
    box-shadow: 0 16px 60px rgba(0,0,0,0.35);
    display: flex;
    flex-direction: column;
  }}

  .json-merger-header {{
    display: flex;
    align-items: center;
    justify-content: space-between;
    padding: 12px 16px;
    border-bottom: 1px solid #e7dfd4;
    background: #f8f4ee;
    color: #3b2f24;
    font-weight: 700;
  }}

  .json-merger-close {{
    border: 1px solid #cbb89b;
    background: #fff;
    color: #3b2f24;
    border-radius: 8px;
    padding: 6px 10px;
    cursor: pointer;
    font-size: 0.95rem;
  }}

  .json-merger-frame {{
    width: 100%;
    height: 100%;
    border: none;
    flex: 1;
  }}

  /* ===== HAMBURGER ===== */
  .menu-toggle {{
    display: none;
    position: fixed;
    top: 15px;
    left: 15px;
    z-index: 200;
    background: var(--sidebar-bg);
    color: var(--sidebar-text);
    border: none;
    border-radius: 8px;
    width: 44px;
    height: 44px;
    font-size: 1.4rem;
    cursor: pointer;
    box-shadow: var(--shadow);
  }}

  /* ===== MAIN CONTENT ===== */
  .main {{
    margin-left: 320px;
    flex: 1;
    padding: 40px;
    display: flex;
    justify-content: center;
  }}

  .book-page {{
    max-width: 780px;
    width: 100%;
  }}

  /* ===== CHAPTER ===== */
  .chapter {{
    background: var(--page-bg);
    border: 1px solid var(--page-border);
    border-radius: 4px;
    padding: 50px 60px;
    margin-bottom: 40px;
    box-shadow: var(--shadow), inset 0 0 80px rgba(0,0,0,0.02);
    position: relative;
  }}

  .chapter::before {{
    content: '';
    position: absolute;
    left: 0;
    top: 0;
    bottom: 0;
    width: 4px;
    background: var(--accent);
    border-radius: 4px 0 0 4px;
  }}

  .chapter-title {{
    font-size: 1.8rem;
    font-weight: 700;
    color: var(--accent);
    text-align: center;
    padding-bottom: 20px;
    margin-bottom: 30px;
    border-bottom: 2px solid var(--accent-light);
    line-height: 1.4;
  }}

  /* ===== SECTION ===== */
  .section {{
    margin-top: 35px;
    padding-top: 25px;
    border-top: 1px dashed var(--page-border);
  }}

  .section-title {{
    font-size: 1.3rem;
    font-weight: 600;
    color: var(--accent);
    margin-bottom: 18px;
    padding-left: 14px;
    border-left: 3px solid var(--accent-light);
    line-height: 1.4;
  }}

  /* ===== CONTENT BLOCKS ===== */
  .bengali-text {{
    font-size: 1.05rem;
    line-height: 2;
    margin-bottom: 16px;
    text-align: justify;
    color: var(--text-dark);
  }}

  .arabic-text {{
    direction: rtl;
    font-family: 'Noto Naskh Arabic', 'Traditional Arabic', serif;
    font-size: 1.25rem;
    line-height: 2.2;
    background: var(--arabic-bg);
    border: 1px solid var(--arabic-border);
    border-radius: 6px;
    padding: 18px 24px;
    margin: 20px 0;
    text-align: right;
    color: #1a1200;
  }}

  /* ===== PAGE MARKERS ===== */
  .page-marker {{
    text-align: center;
    font-size: 0.78rem;
    color: var(--text-muted);
    margin: 24px 0 8px;
    opacity: 0.7;
    letter-spacing: 0.5px;
  }}

  .page-marker::before,
  .page-marker::after {{
    content: ' \\2014\\2014 ';
    opacity: 0.4;
  }}

  /* ===== SCROLL TO TOP ===== */
  .scroll-top {{
    position: fixed;
    bottom: 30px;
    right: 30px;
    width: 44px;
    height: 44px;
    background: var(--accent);
    color: #fff;
    border: none;
    border-radius: 50%;
    font-size: 1.2rem;
    cursor: pointer;
    box-shadow: var(--shadow);
    opacity: 0;
    transition: opacity 0.3s;
    z-index: 50;
  }}

  .scroll-top.visible {{
    opacity: 1;
  }}

  /* Loading State */
  .loading {{
    text-align: center;
    padding: 60px;
    color: var(--text-muted);
    font-size: 1.2rem;
  }}

  /* ===== RESPONSIVE ===== */
  @media (max-width: 900px) {{
    .sidebar {{
      transform: translateX(-100%);
    }}
    .sidebar.open {{
      transform: translateX(0);
    }}
    .menu-toggle {{
      display: block;
    }}
    .main {{
      margin-left: 0;
      padding: 20px;
      padding-top: 70px;
    }}
    .chapter {{
      padding: 30px 25px;
    }}
  }}

  /* ===== PRINT ===== */
  @media print {{
    .sidebar, .menu-toggle, .scroll-top, .json-merger-modal {{ display: none; }}
    .main {{ margin-left: 0; padding: 0; }}
    .chapter {{ box-shadow: none; border: none; page-break-inside: avoid; }}
    .section {{ page-break-inside: avoid; }}
  }}
</style>
</head>
<body>

<button class="menu-toggle" id="menuToggle">&#9776;</button>

<nav class="sidebar" id="sidebar">
  <div class="sidebar-header">
    <h2>Islamic Books</h2>
    <select class="book-selector" id="bookSelector">
      {book_options}
    </select>
    <p class="book-info" id="bookInfo"></p>
  </div>
  <ul class="toc" id="toc"></ul>
  <div class="sidebar-footer">
    <button class="sidebar-action" id="jsonMergerBtn" type="button">
      JSON Merger
    </button>
  </div>
</nav>

<main class="main">
  <div class="book-page" id="bookContent">
    <div class="loading">Loading...</div>
  </div>
</main>

<button class="scroll-top" id="scrollTop">&#8593;</button>

<div class="json-merger-modal" id="jsonMergerModal" aria-hidden="true">
  <div class="json-merger-dialog" role="dialog" aria-modal="true" aria-label="JSON Merger">
    <div class="json-merger-header">
      <span>JSON Merger</span>
      <button class="json-merger-close" id="jsonMergerCloseBtn" type="button">Close</button>
    </div>
    <iframe class="json-merger-frame" src="archive/merge_json_files.html" title="JSON Merger"></iframe>
  </div>
</div>

<script>
// Books index (metadata only - individual books loaded on demand)
const booksIndex = {index_json};

// Cache for loaded books
const booksCache = {{}};

// DOM Elements
const bookSelector = document.getElementById('bookSelector');
const bookInfo = document.getElementById('bookInfo');
const toc = document.getElementById('toc');
const bookContent = document.getElementById('bookContent');
const sidebar = document.getElementById('sidebar');
const menuToggle = document.getElementById('menuToggle');
const scrollBtn = document.getElementById('scrollTop');
const jsonMergerBtn = document.getElementById('jsonMergerBtn');
const jsonMergerModal = document.getElementById('jsonMergerModal');
const jsonMergerCloseBtn = document.getElementById('jsonMergerCloseBtn');

// Current book
let currentBook = null;

// Parse tagged text to HTML
function parseTaggedText(text) {{
  if (!text) return '';

  const parts = text.split(/(<[ca]>.*?<\\/[ca]>)/gs);
  let html = '';

  for (const part of parts) {{
    const cMatch = part.match(/<c>(.*?)<\\/c>/s);
    const aMatch = part.match(/<a>(.*?)<\\/a>/s);

    if (cMatch) {{
      const lines = cMatch[1].trim().split('\\n').filter(l => l.trim());
      html += `<div class="bengali-text">${{lines.join('<br>')}}</div>`;
    }} else if (aMatch) {{
      const lines = aMatch[1].trim().split('\\n').filter(l => l.trim());
      html += `<div class="arabic-text">${{lines.join('<br>')}}</div>`;
    }}
  }}

  return html;
}}

// Build content HTML
function buildContentHtml(contentList) {{
  let html = '';
  for (const item of contentList) {{
    if (item.page) {{
      html += `<div class="page-marker">${{item.page}}</div>`;
    }}
    html += parseTaggedText(item.text);
  }}
  return html;
}}

// Load book data from JSON file
async function loadBook(bookId) {{
  // Return from cache if already loaded
  if (booksCache[bookId]) {{
    return booksCache[bookId];
  }}

  // Show loading state
  bookContent.innerHTML = '<div class="loading">Loading book...</div>';

  try {{
    const response = await fetch(`output/${{bookId}}.json`);
    if (!response.ok) {{
      throw new Error(`Failed to load book: ${{response.status}}`);
    }}
    const bookData = await response.json();

    // Store in cache
    booksCache[bookId] = bookData;

    return bookData;
  }} catch (error) {{
    console.error('Error loading book:', error);
    bookContent.innerHTML = `<div class="loading">Error loading book: ${{error.message}}</div>`;
    return null;
  }}
}}

// Render book
async function renderBook(bookId) {{
  const book = await loadBook(bookId);
  if (!book) return;

  currentBook = book;

  // Update info
  bookInfo.textContent = `${{book.chapters.length}} chapters`;

  // Build TOC
  let tocHtml = '';
  for (const ch of book.chapters) {{
    tocHtml += `<li class="toc-chapter"><a href="#ch-${{ch.id}}">${{ch.title}}</a>`;
    if (ch.sections && ch.sections.length > 0) {{
      tocHtml += '<ul class="toc-sections">';
      for (const sec of ch.sections) {{
        tocHtml += `<li><a href="#sec-${{sec.id}}">${{sec.title}}</a></li>`;
      }}
      tocHtml += '</ul>';
    }}
    tocHtml += '</li>';
  }}
  toc.innerHTML = tocHtml;

  // Build body
  let bodyHtml = '';
  for (const ch of book.chapters) {{
    bodyHtml += `<div class="chapter" id="ch-${{ch.id}}">`;
    bodyHtml += `<h1 class="chapter-title">${{ch.title}}</h1>`;

    if (ch.intro_content && ch.intro_content.length > 0) {{
      bodyHtml += '<div class="chapter-intro">';
      bodyHtml += buildContentHtml(ch.intro_content);
      bodyHtml += '</div>';
    }}

    for (const sec of ch.sections || []) {{
      bodyHtml += `<div class="section" id="sec-${{sec.id}}">`;
      bodyHtml += `<h2 class="section-title">${{sec.title}}</h2>`;
      bodyHtml += buildContentHtml(sec.content || []);
      bodyHtml += '</div>';
    }}

    bodyHtml += '</div>';
  }}
  bookContent.innerHTML = bodyHtml;

  // Re-attach TOC event listeners
  attachTocListeners();
  attachObserver();

  // Scroll to top
  window.scrollTo({{ top: 0, behavior: 'smooth' }});
}}

// TOC event listeners
function attachTocListeners() {{
  document.querySelectorAll('.toc-chapter > a').forEach(link => {{
    link.addEventListener('click', (e) => {{
      const li = link.parentElement;
      document.querySelectorAll('.toc-chapter').forEach(el => {{
        if (el !== li) el.classList.remove('expanded');
      }});
      li.classList.toggle('expanded');
    }});
  }});

  sidebar.querySelectorAll('a').forEach(a => {{
    a.addEventListener('click', () => {{
      if (window.innerWidth <= 900) sidebar.classList.remove('open');
    }});
  }});
}}

// Intersection observer for active section
function attachObserver() {{
  const observer = new IntersectionObserver((entries) => {{
    entries.forEach(entry => {{
      if (entry.isIntersecting) {{
        const id = entry.target.id;
        document.querySelectorAll('.sidebar a').forEach(a => a.classList.remove('active'));
        const active = document.querySelector(`.sidebar a[href="#${{id}}"]`);
        if (active) {{
          active.classList.add('active');
          const parentLi = active.closest('.toc-chapter');
          if (parentLi) parentLi.classList.add('expanded');
        }}
      }}
    }});
  }}, {{ threshold: 0.2 }});

  document.querySelectorAll('.chapter, .section').forEach(el => observer.observe(el));
}}

function openJsonMergerModal() {{
  jsonMergerModal.classList.add('open');
  jsonMergerModal.setAttribute('aria-hidden', 'false');
  document.body.style.overflow = 'hidden';
  if (window.innerWidth <= 900) sidebar.classList.remove('open');
}}

function closeJsonMergerModal() {{
  jsonMergerModal.classList.remove('open');
  jsonMergerModal.setAttribute('aria-hidden', 'true');
  document.body.style.overflow = '';
}}

// Event Listeners
bookSelector.addEventListener('change', (e) => {{
  renderBook(e.target.value);
}});

menuToggle.addEventListener('click', () => sidebar.classList.toggle('open'));

window.addEventListener('scroll', () => {{
  scrollBtn.classList.toggle('visible', window.scrollY > 400);
}});

scrollBtn.addEventListener('click', () => {{
  window.scrollTo({{ top: 0, behavior: 'smooth' }});
}});

jsonMergerBtn.addEventListener('click', openJsonMergerModal);
jsonMergerCloseBtn.addEventListener('click', closeJsonMergerModal);

jsonMergerModal.addEventListener('click', (e) => {{
  if (e.target === jsonMergerModal) {{
    closeJsonMergerModal();
  }}
}});

window.addEventListener('keydown', (e) => {{
  if (e.key === 'Escape' && jsonMergerModal.classList.contains('open')) {{
    closeJsonMergerModal();
  }}
}});

// Initial render
if (booksIndex.length > 0) {{
  renderBook(booksIndex[0].id);
}}
</script>
</body>
</html>'''

    return html

# ==========================================
# MAIN EXECUTION
# ==========================================

def main():
    print("=" * 50)
    print("Multi-Book Processor")
    print("=" * 50)

    # Get all txt files in books directory
    books_dir = Path(BOOKS_DIR)
    hierarchy_dir = Path(HIERARCHY_DIR)

    if not books_dir.exists():
        print(f"Error: '{BOOKS_DIR}' directory not found!")
        print(f"Please create the directory and add your .txt files.")
        return

    txt_files = list(books_dir.glob("*.txt"))

    if not txt_files:
        print(f"No .txt files found in '{BOOKS_DIR}' directory!")
        return

    print(f"\nFound {len(txt_files)} book(s) to process:")
    for f in txt_files:
        print(f"  - {f.name}")

    all_books_data = []

    for txt_file in txt_files:
        book_name = txt_file.stem  # filename without extension
        hierarchy_file = hierarchy_dir / f"{book_name}.json"

        if not hierarchy_file.exists():
            print(f"\n  Warning: No hierarchy file found for '{book_name}'")
            print(f"  Expected: {hierarchy_file}")
            print(f"  Skipping this book...")
            continue

        try:
            book_data = process_book(book_name, txt_file, hierarchy_file)

            all_books_data.append({
                "id": book_name.lower().replace(" ", "_").replace("-", "_"),
                "name": book_name,
                "chapters": book_data
            })

            print(f"  Done! ({len(book_data)} chapters)")

        except Exception as e:
            print(f"  Error processing '{book_name}': {e}")
            import traceback
            traceback.print_exc()

    if not all_books_data:
        print("\nNo books were processed successfully!")
        return

    # Save INDIVIDUAL JSON files for each book (prevents huge single file)
    # Also save XLSX files in "File Sheet" folder
    books_index = []
    for book in all_books_data:
        book_id = book["id"]
        book_json_path = Path(OUTPUT_DIR) / f"{book_id}.json"
        book_xlsx_path = Path(XLSX_DIR) / f"{book_id}.xlsx"

        # Save individual book data as JSON
        with open(book_json_path, "w", encoding="utf-8") as f:
            json.dump(book, f, ensure_ascii=False, indent=2)
        print(f"  Saved JSON: {book_json_path}")

        # Save individual book data as XLSX
        save_book_as_xlsx(book, book["name"], book_xlsx_path)

        # Add to index (metadata only, no content)
        books_index.append({
            "id": book_id,
            "name": book["name"],
            "chapter_count": len(book["chapters"])
        })

    # Save books index file (small file with list of available books)
    index_path = Path(OUTPUT_DIR) / "books_index.json"
    with open(index_path, "w", encoding="utf-8") as f:
        json.dump(books_index, f, ensure_ascii=False, indent=2)
    print(f"\nSaved books index to: {index_path}")

    # Generate HTML viewer (loads books dynamically)
    print("\nGenerating HTML viewer...")
    html_content = generate_multi_book_viewer(books_index)

    with open("book_viewer.html", "w", encoding="utf-8") as f:
        f.write(html_content)

    print("=" * 50)
    print("SUCCESS!")
    print("=" * 50)
    print(f"Processed {len(all_books_data)} book(s)")
    print(f"JSON files saved in: {OUTPUT_DIR}/")
    print(f"XLSX files saved in: {XLSX_DIR}/")
    print(f"HTML viewer: book_viewer.html")
    print("\nOpen book_viewer.html in any browser to view your books!")

if __name__ == "__main__":
    main()
